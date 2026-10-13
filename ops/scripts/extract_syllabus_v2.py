"""Extractor conservador: formularios individuales; comparativas a revisión."""
import argparse
import json
import re
import unicodedata
from pathlib import Path
import openpyxl

PATTERNS = {'anio': ['ano'], 'codigo_plan_estudios': ['codigo plan de estudio'], 'codigo_unidad_estudio': ['codigo unidad de estudio'], 'tipo_unidad_estudio': ['tipo unidad de estudio'], 'facultad_area': ['facultad o area'], 'nivel': ['nivel'], 'programa': ['programa'], 'nombre_unidad_estudio': ['nombre unidad de estudio', 'nombre unidad  de estudio'], 'creditos_academicos': ['creditos academicos'], 'had': ['had'], 'hta': ['hta'], 'ht': ['ht'], 'modalidad': ['modalidad'], 'descripcion_unidad': ['descripcion de la unidad de estudio', 'proposito de la unidad de estudio'], 'resultado_aprendizaje_unidad': ['resultado de aprendizaje de la unidad de estudio', 'resultado de aprendizaje unidad de estudio', 'resultado de aprendizaje'], 'contenidos_unidad': ['contenidos unidad de estudio', 'contenidos de la unidad de estudio'], 'esquema_calificacion': ['esquema de calificacion'], 'recursos_tecnologicos': ['recursos tecnologicos'], 'componente_lengua_extranjera': ['componente de lengua extranjera'], 'bibliografia_obligatoria': ['bibliografia obligatoria'], 'bibliografia_complementaria': ['bibliografia complementaria'], 'prerrequisito': ['prerrequisito', 'pre requisito', 'requisito previo'], 'competencia_programa': ['competencia de programa'], 'estrategias_pedagogicas': ['estrategias pedagogicas'], 'modulo': ['modulo'], 'porcentaje': ['porcentaje'], 'ciclo': ['ciclo'], 'semestre': ['semestre'], 'control_cambios': ['control de cambios'], 'resultado_aprendizaje_programa': ['resultado de aprendizaje de programa'], 'criterios_evaluacion': ['criterios de evaluacion'], 'competencias_modelo': ['competencias del modelo educativo']}
BLOCK = {'descripcion_unidad', 'resultado_aprendizaje_unidad',
         'contenidos_unidad', 'competencia_programa', 'esquema_calificacion',
         'bibliografia_obligatoria', 'bibliografia_complementaria',
         'recursos_tecnologicos', 'estrategias_pedagogicas'}
MINOR = {'modulo', 'ciclo', 'semestre'}

def norm(v):
    s = unicodedata.normalize('NFKD', str(v or ''))
    s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r'\s+', ' ', s).strip().lower()

def field(v):
    t = re.sub(r'^\d+[.\s]+', '', norm(v))
    matches = [(len(p), k) for k, ps in PATTERNS.items() for p in ps
               if t.startswith(p)
               and (k not in MINOR or t == p or t.startswith(p + ":"))
               and (p != "programa" or re.match(r"^programas?(?:\b|:)", t))]
    if not matches:
        return None
    k = max(matches)[1]
    if k == 'resultado_aprendizaje_unidad' and 'programa' in t:
        return 'resultado_aprendizaje_programa'
    if k.startswith('bibliografia') and 'obligatoria y complementaria' in t:
        return None
    return k

def extract(ws):
    # Tablas comparativas: varias unidades en la misma fila de nombres.
    names = [(c.row, c.column) for row in ws.iter_rows(max_col=min(ws.max_column, 100)) for c in row
             if field(c.value) == 'nombre_unidad_estudio']
    if not names:
        return None, 'sin_etiqueta_nombre'
    r, col = names[0]
    values = [c for c in ws[r] if col < c.column <= (9 if col == 7 else ws.max_column) and c.value is not None and str(c.value).strip()
              and not field(c.value)]
    if len(values) > 1:
        return None, 'tabla_comparativa_requiere_revision'
    labels = []
    for row in ws.iter_rows(max_col=min(ws.max_column, 100)):
        for c in row:
            k = field(c.value)
            if k:
                labels.append((c.row, c.column, k, str(c.value)))
    record, cells = {}, {}
    labels.sort(key=lambda item: (bool(re.match(r"^\s*\d+\s*\.", item[3])), item[0], item[1]))
    for r, col, k, raw in labels:
        if k in record or k in MINOR or k == 'control_cambios':
            continue
        end = min([rr for rr, cc, kk, vv in labels
                   if rr > r and kk not in MINOR] or [ws.max_row + 1]) - 1
        # Encabezados numerados de secciones también delimitan bloques.
        section_rows = [c.row for row in ws.iter_rows(max_col=min(ws.max_column, 100)) for c in row
                        if c.row > r and c.column <= col and isinstance(c.value, str)
                        and re.match(r'^\s*\d+\s*\.\s+[A-Za-zÁÉÍÓÚáéíóú]', c.value)]
        if section_rows:
            end = min(end, min(section_rows) - 1)
        right = min([cc for rr, cc, kk, vv in labels
                     if rr == r and cc > col] or [(10 if col <= 9 and names[0][1] == 7 else min(ws.max_column, 100) + 1)]) - 1
        parts, refs = [], []
        if ':' in raw and k.startswith('bibliografia'):
            tail = raw.split(':', 1)[1].strip()
            if tail:
                parts.append(tail)
                refs.append(ws.cell(r, col).coordinate)
        candidates = [ws.cell(r, c) for c in range(col + 1, right + 1)]
        if k in BLOCK:
            candidates += [ws.cell(rr, cc) for rr in range(r + 1, end + 1)
                           for cc in range(col, right + 1)]
        for c in candidates:
            if c.value is None or field(c.value):
                continue
            value = str(c.value).strip()
            if value and value not in parts:
                parts.append(value)
                refs.append(c.coordinate)
            if parts and k not in BLOCK:
                break
        record[k] = '\n'.join(parts) or None
        cells[k] = refs
    if not record.get('nombre_unidad_estudio') or not any(record.get(k) for k in ['descripcion_unidad', 'resultado_aprendizaje_unidad', 'contenidos_unidad']):
        return None, 'formulario_incompleto'
    # Separar criterios incrustados sin descartar su contenido.
    value = record.get('resultado_aprendizaje_unidad')
    if value:
        match = re.search(r'(?im)^\s*criterios de evaluaci[oó]n\s*:', value)
        if match:
            record['criterios_evaluacion_integrados'] = value[match.end():].strip()
            record['resultado_aprendizaje_unidad'] = value[:match.start()].strip() or None
    if norm(record.get('descripcion_unidad')) in {'descripcion delcurso', 'descripcion del curso'}:
        return None, 'plantilla_con_instrucciones'
    record['requiere_revision_identidad'] = not bool(record.get('codigo_unidad_estudio'))
    record['celdas_origen'] = cells
    return record, None


def extract_row_table(ws):
    for row in ws.iter_rows(max_row=min(ws.max_row, 10), max_col=min(ws.max_column, 100)):
        mapping = {field(c.value): c.column for c in row if field(c.value)}
        if not {'nombre_unidad_estudio', 'codigo_unidad_estudio', 'codigo_plan_estudios'} <= mapping.keys():
            continue
        if len(mapping) < 8:
            continue
        records = []
        for rr in range(row[0].row + 1, ws.max_row + 1):
            record = {k: str(ws.cell(rr, cc).value).strip() if ws.cell(rr, cc).value is not None else None
                      for k, cc in mapping.items()}
            if not record.get('nombre_unidad_estudio'):
                continue
            record['fila_origen'] = rr
            record['celdas_origen'] = {k: [ws.cell(rr, cc).coordinate] for k, cc in mapping.items()}
            record['requiere_revision_identidad'] = not bool(record.get('codigo_unidad_estudio'))
            records.append(record)
        return records
    return None

def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('root', type=Path)
    ap.add_argument('--out', type=Path, required=True)
    args = ap.parse_args()
    if not args.root.is_dir():
        ap.error('root debe ser una carpeta')
    records, audit = [], []
    for path in sorted(args.root.rglob('*.xlsx')):
        if path.name.startswith('~$'):
            continue
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            for ws in wb:
                table = extract_row_table(ws)
                if table is not None:
                    audit.append({'archivo': str(path), 'hoja': ws.title,
                                  'estado': 'tabla_por_filas_extraida', 'registros': len(table)})
                    for item in table:
                        item.update(file_path=str(path), folder_breadcrumb=str(path.parent.relative_to(args.root)),
                                    sheet_name=ws.title, sheet_type='master_table_row', extractor_version='syllabus_v2_3')
                        records.append(item)
                    continue
                rec, reason = extract(ws)
                audit.append({'archivo': str(path), 'hoja': ws.title,
                              'estado': reason or 'extraida'})
                if rec:
                    rec.update(file_path=str(path),
                               folder_breadcrumb=str(path.parent.relative_to(args.root)),
                               sheet_name=ws.title, sheet_type='unit_detail',
                               extractor_version='syllabus_v2_3')
                    records.append(rec)
            wb.close()
        except Exception as exc:
            audit.append({'archivo': str(path), 'estado': 'error', 'detalle': str(exc)})
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding='utf-8')
    audit_path = args.out.with_name(args.out.stem + '_auditoria.json')
    audit_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'Registros extraídos: {len(records)}')
    print(f'JSON: {args.out}')
    print(f'Auditoría: {audit_path}')
    print('No se modificó Neo4j. Revisar auditoría antes de sincronizar.')

if __name__ == '__main__':
    main()
